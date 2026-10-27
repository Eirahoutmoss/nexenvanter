"""Uçtan uca çekirdek testleri: python -m pytest -q"""
import io
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

H = {"x-nx": "1"}


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    from app.paths import DIRS
    DIRS.init(str(tmp_path_factory.mktemp("data")))
    from app import api, db
    db.init()
    c = TestClient(api.app)
    r = c.post("/api/setup", json={"username": "admin", "password": "Parola123", "full_name": "Yönetici"}, headers=H)
    assert r.status_code == 200, r.text
    return c


def _xlsx(rows):
    wb = openpyxl.Workbook()
    for r in rows:
        wb.active.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def test_textnorm():
    from app.textnorm import fold, header_key, norm_date, norm_mac, tr_upper
    assert fold("İSTANBUL Işık") == "istanbul isik"
    assert header_key("BULUNDUĞU BİRİM \nSERVİS\nALAN") == "bulundugu birim servis alan"
    assert norm_mac("aa-bb-cc-dd-ee-ff") == "AA:BB:CC:DD:EE:FF"
    assert norm_date("15.03.2024") == "2024-03-15"
    assert norm_date(45366) == "2024-03-15"
    assert tr_upper("bilgi işlem") == "BİLGİ İŞLEM"
    with pytest.raises(ValueError):
        norm_mac("12:34")


def test_csrf_and_auth(client):
    assert client.post("/api/assets", json={"asset_code": "X1"}).status_code == 403
    anon = TestClient(client.app)
    assert anon.get("/api/assets").status_code == 401


def test_crud_and_duplicate(client):
    r = client.post("/api/assets", json={"asset_code": "0001245", "brand_model": "Fujitsu Esprimo D556",
                                         "location_name": "Bilgi İşlem", "mac_address": "aabbccddeeff"}, headers=H)
    assert r.status_code == 200, r.text
    a = r.json()
    assert a["asset_code"] == "0001245" and a["data"]["mac_address"] == "AA:BB:CC:DD:EE:FF"
    assert client.post("/api/assets", json={"asset_code": "0001245"}, headers=H).status_code == 409
    r = client.patch(f"/api/assets/{a['id']}", json={"status": "bakımda"}, headers=H)
    assert r.json()["status"] == "bakımda"
    assert client.get("/api/assets", params={"q": "bilgi islem"}).json()["total"] == 1
    d = client.get(f"/api/assets/{a['id']}").json()
    assert [e["action"] for e in d["events"]] == ["update", "create"]
    p = TestClient(client.app).get(f"/e/{a['public_id']}")
    assert p.status_code == 200 and "Fujitsu" in p.text and "AA:BB" not in p.text  # MAC herkese açık değil


def test_import_real_file(client):
    src = Path(__file__).parent / "fixtures" / "test.xlsx"
    r = client.post("/api/import/upload", files={"file": ("test.xlsx", src.read_bytes())}, headers=H)
    assert r.status_code == 200, r.text
    an = r.json()
    sug = {c["header"]: c["suggestion"] for c in an["columns"]}
    assert sug["NO"] == "__ignore__" and sug["BARKOT NUMARASI"] == "asset_code"
    assert sug["BULUNDUĞU BİRİM SERVİS ALAN"] == "location_name" and sug["MAC ADRES"] == "mac_address"
    mapping = {str(c["index"]): c["suggestion"] for c in an["columns"]}
    body = {"token": an["token"], "sheet": an["sheet"], "header_row": an["header_row"], "mapping": mapping}
    dry = client.post("/api/import/run", json={**body, "dry_run": True}, headers=H).json()
    assert dry["stats"]["inserted"] == 30
    assert client.get("/api/assets", params={"q": "1000001"}).json()["total"] == 0  # dry-run yazmaz
    real = client.post("/api/import/run", json={**body, "dry_run": False}, headers=H).json()
    assert real["stats"]["inserted"] == 30
    again = client.post("/api/import/run", json={**body, "dry_run": False}, headers=H).json()
    assert again["stats"]["skipped"] == 30
    a = client.get("/api/assets/by-code/1000001").json()
    assert a["data"]["serial_number"] == "123123123" and a["data"]["brand_model"] == "FUJİTSU"


def test_import_new_column_creates_field(client):
    raw = _xlsx([["BARKOD NO", "MARKA MODEL", "ODA NO", "KURULUM TARİHİ"],
                 ["2000001", "HP", "10.0.0.5", "01.01.2027"], ["2000001", "HP", "", ""], ["", "boş barkod", "", ""]])
    an = client.post("/api/import/upload", files={"file": ("yeni.xlsx", raw)}, headers=H).json()
    mapping = {str(c["index"]): c["suggestion"] for c in an["columns"]}
    assert mapping["2"] == "__new__" and mapping["3"] == "__new__"
    res = client.post("/api/import/run", json={
        "token": an["token"], "sheet": an["sheet"], "header_row": an["header_row"], "mapping": mapping,
        "new_fields": {"2": {"label": "Oda No"}, "3": {"label": "Kurulum Tarihi", "type": "date"}},
        "dry_run": False}, headers=H).json()
    st = res["stats"]
    assert st["inserted"] == 1 and st["errors"] == 2 and st["new_fields"] == ["Oda No", "Kurulum Tarihi"]
    a = client.get("/api/assets/by-code/2000001").json()
    assert a["data"]["oda_no"] == "10.0.0.5" and a["data"]["kurulum_tarihi"] == "2027-01-01"
    # sonraki dosyada aynı başlık artık otomatik eşleşir
    an2 = client.post("/api/import/upload", files={"file": ("y2.xlsx", _xlsx([["Barkod", "Oda No"], ["3", "x"]]))},
                      headers=H).json()
    assert an2["columns"][1]["suggestion"] == "oda_no"


def test_labels_and_export(client):
    a = client.get("/api/assets/by-code/1000001").json()
    r = client.get(f"/api/labels/preview/{a['id']}.png")
    assert r.status_code == 200 and r.content[:4] == b"\x89PNG"
    r = client.get("/api/labels/pdf", params={"ids": str(a["id"])})
    assert r.content[:4] == b"%PDF"
    r = client.get("/api/export.xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    assert wb.active.max_row >= 32


def test_roles(client):
    client.post("/api/users", json={"username": "izleyici", "password": "Parola123", "role": "viewer"}, headers=H)
    v = TestClient(client.app)
    assert v.post("/api/login", json={"username": "izleyici", "password": "Parola123"}, headers=H).status_code == 200
    assert v.get("/api/assets").status_code == 200
    assert v.post("/api/assets", json={"asset_code": "Z"}, headers=H).status_code == 403
    assert v.get("/api/users").status_code == 403
