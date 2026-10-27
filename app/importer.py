"""Excel içe aktarma: analiz → eşleştirme → deneme (dry-run) → uygula. Yeni sütunlar yeni alan olur."""
import json
import secrets
import time

import openpyxl

from . import assets as A, db, fields as F
from .paths import DIRS
from .textnorm import cell_text, header_key

MAX_BYTES = 50 * 1024 * 1024
IGNORE, NEW = "__ignore__", "__new__"


class _Rollback(Exception):
    pass


def save_upload(raw: bytes, filename: str) -> str:
    if len(raw) > MAX_BYTES:
        raise ValueError("Dosya 50 MB sınırını aşıyor")
    if not filename.lower().endswith((".xlsx", ".xlsm")) or raw[:2] != b"PK":
        raise ValueError("Yalnızca .xlsx dosyaları desteklenir")
    for p in DIRS.tmp.glob("imp_*.xlsx"):  # 1 günden eski geçici dosyaları temizle
        if p.stat().st_mtime < time.time() - 86400:
            p.unlink(missing_ok=True)
    token = secrets.token_hex(8)
    (DIRS.tmp / f"imp_{token}.xlsx").write_bytes(raw)
    (DIRS.tmp / f"imp_{token}.json").write_text(json.dumps({"filename": filename}), encoding="utf-8")
    return token


def _path(token: str):
    if not token.isalnum():
        raise ValueError("Geçersiz dosya anahtarı")
    p = DIRS.tmp / f"imp_{token}.xlsx"
    if not p.exists():
        raise ValueError("Yüklenen dosya bulunamadı; tekrar yükleyin")
    return p


def _open(token):
    return openpyxl.load_workbook(_path(token), read_only=True, data_only=True)


def _detect_header(rows: list[tuple], flds, ignore) -> int:
    """İlk 20 satırda en çok alanla eşleşen ve en çok dolu metin hücresi olan satır (1 tabanlı)."""
    best, score = 1, -1
    for i, r in enumerate(rows[:20], start=1):
        texts = [c for c in r if isinstance(c, str) and c.strip()]
        if len(texts) < 2:
            continue
        s = sum(1 for c in texts if F.match_header(c, flds, ignore) not in (None, IGNORE)) * 10 + len(texts)
        if s > score:
            best, score = i, s
    return best


def analyze(token: str, sheet: str | None = None, header_row: int | None = None) -> dict:
    wb = _open(token)
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        head = [r for _, r in zip(range(30), ws.iter_rows(values_only=True))]
        flds = F.all_fields()
        ignore = db.get_setting("import_ignore_headers", [])
        hr = header_row or _detect_header(head, flds, ignore)
        hdr = list(head[hr - 1]) if len(head) >= hr else []
        cols, used = [], set()
        for i, h in enumerate(hdr):
            title = cell_text(h) or ""
            sug = F.match_header(title, flds, ignore) if title else IGNORE
            if sug in used and sug != IGNORE:
                sug = NEW  # aynı alana iki sütun eşleşmesin
            if sug is None:
                sug = NEW
            used.add(sug)
            cols.append({"index": i, "header": " ".join(title.split()), "suggestion": sug})
        sample = [[cell_text(c) for c in r] for r in head[hr:hr + 10]]
        return {
            "token": token, "sheets": wb.sheetnames, "sheet": ws.title, "header_row": hr,
            "row_estimate": max(0, (ws.max_row or 0) - hr), "columns": cols, "sample": sample,
            "fields": [{"key": f["key"], "label": f["label"], "type": f["type"]} for f in flds],
        }
    finally:
        wb.close()


def run(token, sheet, header_row: int, mapping: dict, new_fields: dict | None = None, mode="skip",
        blank_overwrites=False, dry_run=True, user="?") -> dict:
    """mapping: {"<sütun index>": alan_key | "__ignore__" | "__new__"}; new_fields: {"<index>": {"label","type"}}
    mode: 'skip' (mevcut barkod atlanır) | 'update' (mevcut kayıt güncellenir)."""
    if mode not in ("skip", "update"):
        raise ValueError("Geçersiz mod")
    keys = [v for v in mapping.values() if v not in (IGNORE, NEW)]
    if len(keys) != len(set(keys)):
        raise ValueError("Bir alan birden fazla sütuna eşlenmiş")
    if "asset_code" not in keys:
        raise ValueError("Barkod / Envanter No sütunu eşlenmeli")
    meta = json.loads((DIRS.tmp / f"imp_{token}.json").read_text(encoding="utf-8"))
    wb = _open(token)
    report, st = [], {"rows": 0, "inserted": 0, "updated": 0, "unchanged": 0, "skipped": 0, "errors": 0, "blank": 0,
                      "new_fields": []}
    job_id = None
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        with db.tx() as t:
            colmap = {}
            for idx, key in mapping.items():
                if key == NEW:
                    nf = (new_fields or {}).get(str(idx)) or {}
                    label = (nf.get("label") or "").strip()
                    if not label:
                        raise ValueError(f"{int(idx) + 1}. sütun için yeni alan adı gerekli")
                    f = F.create(label, nf.get("type") or "text", aliases=[label])
                    st["new_fields"].append(f["label"])
                    key = f["key"]
                if key != IGNORE:
                    colmap[int(idx)] = key
            if not dry_run:
                job_id = t.execute("INSERT INTO import_jobs(at,user,filename,sheet,mode,mapping) VALUES(?,?,?,?,?,?)",
                                   (db.now(), user, meta["filename"], ws.title, mode, json.dumps(mapping))).lastrowid
            seen = {}
            for rno, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
                payload = {k: (row[i] if i < len(row) else None) for i, k in colmap.items()}
                if all(cell_text(v) is None for v in payload.values()):
                    st["blank"] += 1
                    continue
                st["rows"] += 1
                code = cell_text(payload.get("asset_code"))

                def rep(res, msg=""):
                    report.append({"row": rno, "asset_code": code, "result": res, "message": msg})

                if not code:
                    st["errors"] += 1
                    rep("hata", "Barkod boş")
                    continue
                if header_key(code) in seen:
                    st["errors"] += 1
                    rep("hata", f"Dosyada mükerrer barkod (ilk: satır {seen[header_key(code)]})")
                    continue
                seen[header_key(code)] = rno
                try:
                    ex = t.execute("SELECT id FROM assets WHERE asset_code=?", (code,)).fetchone()
                    if ex:
                        if mode == "skip":
                            st["skipped"] += 1
                            rep("atlandı", "Barkod zaten kayıtlı")
                            continue
                        res = A.update(ex[0], payload, user, blank_overwrites=blank_overwrites, t=t, action="import")
                        if res["changed"]:
                            st["updated"] += 1
                            rep("güncellendi")
                        else:
                            st["unchanged"] += 1
                            rep("değişmedi")
                    else:
                        A.create(payload, user, source_import=job_id, t=t)
                        st["inserted"] += 1
                        rep("eklendi")
                except A.AssetError as e:
                    st["errors"] += 1
                    det = e.extra.get("fields")
                    rep("hata", "; ".join(f"{v}" for v in det.values()) if det else e.msg)
            if dry_run:
                raise _Rollback
            t.execute("UPDATE import_jobs SET stats=?, report=? WHERE id=?",
                      (json.dumps(st, ensure_ascii=False), json.dumps(report, ensure_ascii=False), job_id))
    except _Rollback:
        pass
    finally:
        wb.close()
    if not dry_run:
        db.audit(user, "import", job_id, st)
    return {"job_id": job_id, "dry_run": dry_run, "stats": st, "report": report[:2000], "report_total": len(report)}


def jobs(limit=50):
    rows = db.conn().execute("SELECT id,at,user,filename,sheet,mode,stats FROM import_jobs ORDER BY id DESC LIMIT ?", (limit,))
    return [{**dict(r), "stats": json.loads(r["stats"] or "{}")} for r in rows]


def job_report(job_id: int) -> list[dict]:
    r = db.conn().execute("SELECT report FROM import_jobs WHERE id=?", (job_id,)).fetchone()
    return json.loads(r[0] or "[]") if r else []
